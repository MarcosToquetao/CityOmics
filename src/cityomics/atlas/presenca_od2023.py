"""Pessoas na rua com a Pesquisa Origem-Destino 2023 (parecer R5): compatibilidade, réplica e incerteza.

O mesmo pipeline de ``presenca`` (linha do tempo de cada pessoa, minutos andando na origem e no destino,
grade de 100 m, núcleo gaussiano σ = 300 m, soma por setor) aplicado aos microdados da OD 2023 do Metrô
(``data/external/od/OD-2023-zip/Site_190225_PesquisaOD2023.zip``, publicado em 2025-02-12; campo de
ago/2023 a mai/2024; dias úteis; Córrego Alegre / UTM 23S, EPSG:22523, como a 2017).

- **Compatibilidade:** variáveis, códigos, dias da semana, amostra e totais expandidos de 2017 e 2023,
  com o total oficial de viagens da Tabela 17 de cada edição.
- **Leitura da 2023:** a ausência vem gravada como 0 (sem viagem: ``N_VIAG = 0``, coordenadas 0,
  ``FE_VIA = 0``, ``TIPVG = 2``) e ``ANDA_O``/``ANDA_D`` vêm como texto ("NU" nas viagens a pé);
  ``normaliza_2023`` devolve o formato da 2017 (nulos). Sem isso, quem não viajou some da presença.
- **Pedestres:** duas definições de viagem ativa (contada inteira como rua): a pé + bicicleta, como em
  2017 (``pe_bici``), e só a pé (``pe``; a bicicleta entra só com os minutos andando nas pontas).
- **Quantidades do manuscrito:** densidade de pedestres e de presentes por classe relativa à cidade
  (Extended Data Fig. 2c), ρ de Spearman entre a taxa de roubo por km e por pedestre, curva horária do
  risco por pedestre na cidade (hora do pico e múltiplo da média diária), mediana entre classes da
  correlação entre os perfis horários do roubo e dos pedestres e as mudanças de posição no ranking.
- **Incerteza:** bootstrap de domicílios (``ID_DOM``) estratificado pela zona de residência: a amostra
  da OD é de domicílios sorteados dentro de cada zona (a zona é estrato, não conglomerado). Tudo até a
  soma por setor é linear nos pesos, então cada réplica só multiplica o peso de cada segmento pela
  multiplicidade do seu domicílio e refaz a grade (as coordenadas, as horas e a rasterização dos
  setores ficam pré-calculadas). O crime fica fixo: é a incerteza do denominador.

Crime: roubo e furto de celular em dias úteis de 2024–2025, distritos de exploração (D53), como em
``presenca``. Saídas em ``data/processed/atlas/presenca_2023/``; a tabela por setor fica local.
"""

from __future__ import annotations

import hashlib
import json

import numpy as np
import pandas as pd
import polars as pl

from cityomics.atlas import guard
from cityomics.atlas import presenca as P
from cityomics.atlas.attribution import em_fit_multi
from cityomics.atlas.genes_lote import atlas_dir
from cityomics.external import external_dir

OD23_ZIP = "od/OD-2023-zip/Site_190225_PesquisaOD2023.zip"
OD23_SHA256 = "1844d38ad83c86cbb679f24c04e317aa49704c5009f8e5a0d0246f026f06bd9d"
OD23_DBF = "od/OD-2023/Site_190225/Banco2023_divulgacao_190225.dbf"
TAB17 = {"2017": "od/OD-2017/Tabelas-OD2017/Tab17_OD2017.xlsx",
         "2023": "od/OD-2023/Site_190225/Tabelas_Site_OD2023_REV_190225.xlsx"}
COLS = P.COLS + ["ID_DOM", "ZONA", "F_PESS", "DIA_SEM", "MUNI_O", "MUNI_D", "DATA"]
VIAGEM = ["N_VIAG", "FE_VIA", "DIA_SEM", "CO_O_X", "CO_O_Y", "CO_D_X", "CO_D_Y", "H_SAIDA", "MIN_SAIDA", "H_CHEG",
          "MIN_CHEG", "ANDA_O", "ANDA_D", "TIPVG", "MUNI_O", "MUNI_D"]
COORDS = ["CO_DOM_X", "CO_DOM_Y", "CO_O_X", "CO_O_Y", "CO_D_X", "CO_D_Y"]
ATIVOS = {"pe_bici": (3, 4), "pe": (3,)}
N_BOOT = 1000
MADRUGADA = list(P.PERIODOS["madrugada"])  # 0h–5h59
SEED = 20261020
ROUBO = "celular_roubo"
CLASSES_TEXTO = ("Indústria e galpões", "Vertical residencial e misto")


def out_dir():
    d = atlas_dir() / "presenca_2023"
    d.mkdir(parents=True, exist_ok=True)
    return d


# --- leitura ------------------------------------------------------------------------------------

def sha256(path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for bloco in iter(lambda: f.read(1 << 20), b""):
            h.update(bloco)
    return h.hexdigest()


def normaliza_2023(od: pd.DataFrame) -> pd.DataFrame:
    """Converte a codificação da OD 2023 (ausência = 0, ANDA como texto) para a da OD 2017 (nulos)."""
    od = od.copy()
    sem = od["N_VIAG"].eq(0).to_numpy()
    for c in ("ANDA_O", "ANDA_D"):
        if c in od:
            od[c] = pd.to_numeric(od[c], errors="coerce")
    for c in [c for c in VIAGEM if c in od]:
        od[c] = od[c].astype(float).mask(sem)
    for c in [c for c in COORDS if c in od]:
        od[c] = od[c].astype(float).where(od[c] > 0)
    for c in ("ID_PESS", "ID_DOM"):
        if c in od:
            od[c] = od[c].astype(str)
    return od


def _ler(rel: str, cols) -> pd.DataFrame:
    import pyogrio

    path = guard.pre2026(external_dir() / rel)
    return pyogrio.read_dataframe(path, columns=list(cols), read_geometry=False)


def read_od2023(cols=COLS) -> pd.DataFrame:
    return normaliza_2023(_ler(OD23_DBF, cols))


def read_od2017(cols=COLS) -> pd.DataFrame:
    if not (external_dir() / P.OD_DBF).exists():
        P.read_od()  # extrai o zip
    od = _ler(P.OD_DBF, cols)
    for c in ("ID_PESS", "ID_DOM"):
        if c in od:
            od[c] = od[c].astype(str)
    return od


def recorte_municipio(od: pd.DataFrame, muni: int = P.MUNI_SP, campo: str = "MUNI_DOM") -> pd.DataFrame:
    """Linhas cujo município (de domicílio, por padrão) é ``muni`` (São Paulo = 36 nas duas edições)."""
    return od[pd.to_numeric(od[campo], errors="coerce").eq(muni)]


# --- compatibilidade ----------------------------------------------------------------------------

def _total_tab17(rel: str) -> float | None:
    """Total oficial de viagens diárias da RMSP (Tabela 17, linha 'Total RMSP', coluna total)."""
    import openpyxl

    path = guard.pre2026(external_dir() / rel)
    if not path.exists():
        return None
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Tab17"] if "Tab17" in wb.sheetnames else wb.worksheets[0]
    for row in ws.iter_rows(values_only=True):
        if isinstance(row[0], str) and row[0].strip().startswith("Total RMSP"):
            return float(row[5])
    return None


def resumo(od: pd.DataFrame, datas: pd.Series) -> dict:
    """Amostra, cobertura e totais expandidos de uma edição já normalizada."""
    trips = od[od["N_VIAG"].notna()]
    pess = od.drop_duplicates("ID_PESS")
    sp_p, sp_v = recorte_municipio(pess), recorte_municipio(trips)
    dentro = trips[trips["MUNI_O"].eq(P.MUNI_SP) & trips["MUNI_D"].eq(P.MUNI_SP)]
    motor = sp_v[sp_v["TIPVG"].isin([1, 2])]

    def exp(df, tipos=None):
        df = df if tipos is None else df[df["TIPVG"].isin(tipos)]
        return float(df["FE_VIA"].sum())

    sem_viagem = pess["ID_PESS"].isin(trips["ID_PESS"]).eq(False)
    toca_sp = trips[trips["MUNI_O"].eq(P.MUNI_SP) | trips["MUNI_D"].eq(P.MUNI_SP)]
    mad = toca_sp[toca_sp["H_SAIDA"].lt(5)]
    return {
        "registros": int(len(od)), "pessoas": int(pess.shape[0]), "domicilios": int(od["ID_DOM"].nunique()),
        "zonas_com_amostra": int(od["ZONA"].nunique()), "viagens": int(len(trips)),
        "campo": [str(datas.min().date()), str(datas.max().date())],
        "dia_sem_viagens": {str(int(k)): int(v) for k, v in trips["DIA_SEM"].value_counts().sort_index().items()},
        "tipvg_viagens": {str(int(k)): int(v) for k, v in trips["TIPVG"].value_counts().sort_index().items()},
        "coord_x": [float(od["CO_DOM_X"].min()), float(od["CO_DOM_X"].max())],
        "coord_y": [float(od["CO_DOM_Y"].min()), float(od["CO_DOM_Y"].max())],
        "rmsp_viagens_expandidas": exp(trips), "rmsp_a_pe_expandidas": exp(trips, [3]),
        "rmsp_bicicleta_expandidas": exp(trips, [4]),
        "sp_pessoas": int(sp_p.shape[0]), "sp_domicilios": int(recorte_municipio(od)["ID_DOM"].nunique()),
        "sp_pessoas_expandidas": float(sp_p["FE_PESS"].sum()), "sp_viagens": int(len(sp_v)),
        "sp_viagens_expandidas": exp(sp_v), "sp_a_pe_expandidas": exp(sp_v, [3]),
        "sp_bicicleta_expandidas": exp(sp_v, [4]),
        "sp_pessoas_sem_viagem_frac": float(sp_p.loc[sem_viagem.loc[sp_p.index], "FE_PESS"].sum() / sp_p["FE_PESS"].sum()),
        "sp_viagens_por_pessoa": exp(sp_v) / float(sp_p["FE_PESS"].sum()),
        "dentro_sp_a_pe_expandidas": exp(dentro, [3]), "dentro_sp_bicicleta_expandidas": exp(dentro, [4]),
        "viagens_amostra_saida_0_5h_tocando_sp": int(len(mad)),
        "viagens_amostra_saida_0_5h_tocando_sp_ativas": int(mad["TIPVG"].isin([3, 4]).sum()),
        "sp_anda_motorizadas_min": float((motor["ANDA_O"].fillna(0) + motor["ANDA_D"].fillna(0)).mul(motor["FE_VIA"]).sum()
                                         / motor["FE_VIA"].sum()),
    }


def compatibilidade(od17: pd.DataFrame, od23: pd.DataFrame, cols23: list[str], cols17: list[str]) -> dict:
    d17 = pd.to_datetime(od17["DATA"].astype(str).str.zfill(8), format="%d%m%Y", errors="coerce")
    d23 = pd.to_datetime(od23["DATA"], errors="coerce")
    out = {"2017": resumo(od17, d17), "2023": resumo(od23, d23)}
    for ano in ("2017", "2023"):
        out[ano]["tab17_total_rmsp"] = _total_tab17(TAB17[ano])
    out["variaveis_so_2017"] = sorted(set(cols17) - set(cols23))
    out["variaveis_so_2023"] = sorted(set(cols23) - set(cols17))
    out["colunas_do_pipeline_presentes_2023"] = all(c in cols23 for c in P.COLS)
    return out


# --- segmentos (cópia parametrizada de presenca.segments) ---------------------------------------

def segmentos(od: pd.DataFrame, ativos=(3, 4)) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Segmentos de presença e de rua como ``presenca.segments``, com o tipo de viagem ativa como parâmetro
    e a unidade de reamostragem (``_u``, inteiro; 0 se ausente) em cada segmento."""
    od = od.copy()
    if "_u" not in od:
        od["_u"] = 0
    trips = od[od["N_VIAG"].notna() & od["H_SAIDA"].notna() & od["H_CHEG"].notna()].sort_values(["ID_PESS", "N_VIAG"])
    dep = trips["H_SAIDA"] * 60 + trips["MIN_SAIDA"].fillna(0)
    arr = trips["H_CHEG"] * 60 + trips["MIN_CHEG"].fillna(0)
    arr = np.where(arr < dep, arr + 1440, arr)
    trips = trips.assign(dep=dep.to_numpy(), arr=arr)

    g = trips.groupby("ID_PESS", sort=False)
    first, last = g.head(1), g.tail(1)
    pres = [pd.DataFrame({"x": first["CO_O_X"].fillna(first["CO_DOM_X"]), "y": first["CO_O_Y"].fillna(first["CO_DOM_Y"]),
                          "a": 0.0, "b": first["dep"], "w": first["FE_PESS"], "u": first["_u"]})]
    nxt = g["dep"].shift(-1)
    mid = trips[nxt.notna()]
    pres.append(pd.DataFrame({"x": mid["CO_D_X"], "y": mid["CO_D_Y"], "a": mid["arr"], "b": nxt[nxt.notna()],
                              "w": mid["FE_PESS"], "u": mid["_u"]}))
    pres.append(pd.DataFrame({"x": last["CO_D_X"], "y": last["CO_D_Y"], "a": last["arr"], "b": 1440.0,
                              "w": last["FE_PESS"], "u": last["_u"]}))
    sem = od[od["N_VIAG"].isna()].drop_duplicates("ID_PESS")
    pres.append(pd.DataFrame({"x": sem["CO_DOM_X"], "y": sem["CO_DOM_Y"], "a": 0.0, "b": 1440.0, "w": sem["FE_PESS"],
                              "u": sem["_u"]}))
    pres = pd.concat(pres, ignore_index=True)
    pres = pres[(pres["b"] > pres["a"]) & pres["x"].gt(0)]

    ativo = trips["TIPVG"].isin(list(ativos))
    dur = trips["arr"] - trips["dep"]
    ao = np.where(ativo, dur / 2, trips["ANDA_O"].fillna(0))
    ad = np.where(ativo, dur / 2, trips["ANDA_D"].fillna(0))
    rua = pd.concat([
        pd.DataFrame({"x": trips["CO_O_X"], "y": trips["CO_O_Y"], "a": trips["dep"], "b": trips["dep"] + ao,
                      "w": trips["FE_VIA"], "u": trips["_u"]}),
        pd.DataFrame({"x": trips["CO_D_X"], "y": trips["CO_D_Y"], "a": trips["arr"] - ad, "b": trips["arr"],
                      "w": trips["FE_VIA"], "u": trips["_u"]}),
    ], ignore_index=True)
    rua = rua[(rua["b"] > rua["a"]) & rua["x"].gt(0)]
    return pres, rua


# --- grade pré-calculada (o mesmo cálculo de presenca.to_setores, separado para reamostrar) -----

def grade(setores) -> dict:
    from rasterio import features
    from rasterio.transform import from_origin

    gx0, gy1, nx, ny = P._grid(tuple(setores.total_bounds))
    ids = features.rasterize(zip(setores.geometry.values, np.arange(1, len(setores) + 1)), out_shape=(ny, nx),
                             transform=from_origin(gx0, gy1, P.RES, P.RES), fill=0, dtype="int32").ravel()
    rp = setores.geometry.representative_point()
    return {"gx0": gx0, "gy1": gy1, "nx": nx, "ny": ny, "ids": ids, "n": len(setores),
            "rix": np.clip(np.floor((rp.x.to_numpy() - gx0) / P.RES).astype(int), 0, nx - 1),
            "riy": np.clip(np.floor((gy1 - rp.y.to_numpy()) / P.RES).astype(int), 0, ny - 1),
            "area_cells": setores.geometry.area.to_numpy() / P.RES ** 2,
            "n_cells": np.bincount(ids, minlength=len(setores) + 1)[1:]}


def celulas(seg: pd.DataFrame, g: dict) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Célula da grade, pessoas-hora por hora (n × 24, já com o peso) e unidade de cada segmento dentro."""
    from pyproj import Transformer

    tr = Transformer.from_crs(P.OD_CRS, P.CRS, always_xy=True)
    x, y = tr.transform(seg["x"].to_numpy(float), seg["y"].to_numpy(float))
    ix = np.floor((x - g["gx0"]) / P.RES).astype(int)
    iy = np.floor((g["gy1"] - y) / P.RES).astype(int)
    ok = (ix >= 0) & (ix < g["nx"]) & (iy >= 0) & (iy < g["ny"])
    W = P.hour_overlap(seg["a"].to_numpy()[ok], seg["b"].to_numpy()[ok]) * seg["w"].to_numpy()[ok, None]
    return iy[ok] * g["nx"] + ix[ok], W, seg["u"].to_numpy()[ok].astype(int)


def somar(cell: np.ndarray, W: np.ndarray, g: dict) -> np.ndarray:
    """Pessoas-hora por setor × hora (n_setores × 24): grade, núcleo gaussiano e soma por setor."""
    from scipy.ndimage import gaussian_filter

    out = np.zeros((g["n"], 24))
    for h in range(24):
        G = np.bincount(cell, weights=W[:, h], minlength=g["nx"] * g["ny"]).reshape(g["ny"], g["nx"])
        Gs = gaussian_filter(G, P.SIGMA_M / P.RES, mode="constant")
        Gs *= G.sum() / max(Gs.sum(), 1e-12)
        soma = np.bincount(g["ids"], weights=Gs.ravel(), minlength=g["n"] + 1)[1:]
        ponto = Gs[g["riy"], g["rix"]] * g["area_cells"]
        out[:, h] = np.where(g["n_cells"] >= 1, soma, ponto)
    return out


# --- métricas do manuscrito ---------------------------------------------------------------------

def crime_cache(ids: pl.DataFrame, E: np.ndarray, names: list[str]) -> dict:
    """Taxas por km (por período e por hora) e perfis do crime, fixos em todas as réplicas."""
    from cityomics.atlas.fenotipo import PERIODOS as PER_ORDEM
    from cityomics.atlas.fenotipo import _contraste

    n = ids.height
    out = {"per": {}}
    for desfecho in ("celular_roubo", "celular_furto"):
        ev = P._crime(ids, desfecho)
        ev = ev[ev["per"].notna()]
        p = ev["per"].map({q: j for j, q in enumerate(PER_ORDEM)}).to_numpy().astype(int)
        Y = np.bincount(ev["i"].to_numpy() * 4 + p, minlength=n * 4).reshape(n, 4).astype(float)
        out["per"][desfecho] = {"Y": Y, "lam": em_fit_multi(E, Y), "eventos": int(len(ev))}
    ev = P._crime(ids, ROUBO)
    ev = ev[ev["h"].notna()]
    Y24 = np.bincount(ev["i"].to_numpy() * 24 + ev["h"].astype(int).to_numpy(), minlength=n * 24).reshape(n, 24).astype(float)
    _, Pc, w = _contraste(em_fit_multi(E, Y24), E)
    out.update({"Pc": Pc, "w": w, "cidade_crime": Y24.sum(0), "eventos_hora": int(len(ev)), "periodos": PER_ORDEM})
    return out


def _ranking(v: np.ndarray, ok: np.ndarray) -> np.ndarray:
    """Posição (1 = maior) entre as classes válidas; NaN fora delas."""
    r = np.full(len(v), np.nan)
    idx = np.where(ok)[0]
    r[idx[np.argsort(-v[idx], kind="stable")]] = np.arange(1, len(idx) + 1)
    return r


def por_pessoa(E: np.ndarray, names: list[str], D: np.ndarray, Y: np.ndarray, lam: np.ndarray, periodos) -> dict:
    """ρ por pessoa (EM com exposição × densidade do período), Spearman contra λ por km e ranking."""
    from scipy.stats import spearmanr

    peso_km = E.sum(axis=0)
    rho, dens_k = np.zeros_like(lam), np.zeros_like(lam)
    for j, q in enumerate(periodos):
        Dp = D[:, list(P.PERIODOS[q])].mean(axis=1)
        Dp = Dp / Dp[Dp > 0].mean()
        rho[:, j] = em_fit_multi(E * Dp[:, None], Y[:, [j]])[:, 0]
        dens_k[:, j] = (E * Dp[:, None]).sum(axis=0) / np.where(peso_km > 0, peso_km, np.nan)
    tot_l, tot_r = lam.sum(axis=1), rho.sum(axis=1)
    km = np.array([not nm.startswith("n_") for nm in names])
    ok = (tot_l > 0) & (tot_r > 0) & km
    fm = np.nansum(dens_k * lam, axis=1) / np.where(tot_l > 0, tot_l, np.nan)
    r_km, r_p = _ranking(tot_l, ok), _ranking(tot_r, ok)
    return {"spearman": float(spearmanr(tot_l[ok], tot_r[ok]).statistic), "rho": rho, "dens_k": dens_k,
            "fator_multidao": fm, "subida": r_km - r_p, "ok": ok,
            "amplitude_log10_multidao": float(np.log10(np.nanmax(fm[ok]) / np.nanmin(fm[ok]))),
            "amplitude_log10_km": float(np.log10(tot_l[ok].max() / tot_l[ok].min()))}


def relogio(E: np.ndarray, names: list[str], D: np.ndarray, R: np.ndarray, cr: dict) -> dict:
    """Curva horária do risco por pedestre na cidade e correlação dos perfis horários por classe."""
    from cityomics.atlas.fenotipo import PESO_MIN

    fr = E.T @ D
    Pr = fr / np.where(fr.sum(1, keepdims=True) > 0, fr.sum(1, keepdims=True), 1.0)
    cor = {nm: float(np.corrcoef(cr["Pc"][k], Pr[k])[0, 1]) for k, nm in enumerate(names) if cr["w"][k] >= PESO_MIN}
    cc = cr["cidade_crime"] / cr["cidade_crime"].sum()
    rd, rp = D.sum(0), R.sum(0)
    risco = cc / (rd / rd.sum())          # definição de presenca._relogio_vs_rua: soma das densidades
    risco_pessoas = cc / (rp / rp.sum())  # variante: soma das pessoas na rua dos setores
    mad = float(cc[MADRUGADA].sum() / (rd[MADRUGADA].sum() / rd.sum()))  # o bloco 0–5h, menos ruidoso que a hora
    return {"risco": risco, "risco_pessoas": risco_pessoas, "pico": int(np.argmax(risco)), "multiplo": float(risco.max()),
            "risco_madrugada": mad,
            "pico_pessoas": int(np.argmax(risco_pessoas)), "multiplo_pessoas": float(risco_pessoas.max()),
            "cor": cor, "cor_mediana": float(np.median(list(cor.values()))), "Pr": Pr,
            "perfil_rua": rd / rd.sum(), "perfil_rua_pessoas": rp / rp.sum()}


def mapa_setores(setor_ids: np.ndarray, ids: pl.DataFrame) -> dict:
    """Posição de cada setor de exploração na ordem de ``load_setores`` e sua área (km²)."""
    from cityomics.atlas.spots import setores_attrs

    area = ids.join(setores_attrs().select("cd_setor", "area_km2"), on="cd_setor", how="left")["area_km2"].fill_null(0).to_numpy()
    i = pd.Series(np.arange(len(setor_ids)), index=setor_ids).reindex(ids["cd_setor"].to_list()).to_numpy()
    m = ~np.isnan(i)
    return {"m": m, "i": i[m].astype(int), "area": area, "n": ids.height}


def densidade(R: np.ndarray, mp: dict) -> tuple[np.ndarray, np.ndarray]:
    """Pessoas por km² (como ``presenca._densidade`` + ``nan_to_num``) e pessoas nos setores de exploração."""
    Rx = np.zeros((mp["n"], 24))
    Rx[mp["m"]] = R[mp["i"]]
    D = np.nan_to_num(Rx / np.where(mp["area"] > 0, mp["area"], np.nan)[:, None])
    return D, Rx


def metricas(E, names, D, R, cr, Dpres=None) -> dict:
    rb = cr["per"][ROUBO]
    pp = por_pessoa(E, names, D, rb["Y"], rb["lam"], cr["periodos"])
    out = {"pessoa": pp, "relogio": relogio(E, names, D, R, cr)}
    if Dpres is not None:
        out["presenca"] = por_pessoa(E, names, Dpres, rb["Y"], rb["lam"], cr["periodos"])
        fu = cr["per"]["celular_furto"]
        out["furto_rua"] = por_pessoa(E, names, D, fu["Y"], fu["lam"], cr["periodos"])["spearman"]
    return out


def _resumo_metricas(m: dict, names: list[str]) -> dict:
    pp, rl = m["pessoa"], m["relogio"]
    ok = pp["ok"]
    out = {"spearman_km_vs_pedestre": pp["spearman"],
           "fator_multidao_rua": {nm: float(pp["fator_multidao"][k]) for k, nm in enumerate(names) if np.isfinite(pp["fator_multidao"][k])},
           "subida_no_ranking": {nm: float(pp["subida"][k]) for k, nm in enumerate(names) if ok[k]},
           "amplitude_log10_multidao": pp["amplitude_log10_multidao"], "amplitude_log10_km": pp["amplitude_log10_km"],
           "hora_pico": rl["pico"], "multiplo_pico": rl["multiplo"], "risco_relativo_por_hora": rl["risco"].tolist(),
           "hora_pico_pessoas": rl["pico_pessoas"], "multiplo_pico_pessoas": rl["multiplo_pessoas"],
           "risco_madrugada_0_5h": rl["risco_madrugada"], "risco_relativo_por_hora_pessoas": rl["risco_pessoas"].tolist(),
           "correlacao_mediana": rl["cor_mediana"], "correlacao_por_classe": rl["cor"],
           "perfil_rua_cidade": rl["perfil_rua"].tolist(), "perfil_rua_cidade_pessoas": rl["perfil_rua_pessoas"].tolist()}
    if "presenca" in m:
        fp = m["presenca"]["fator_multidao"]
        out["fator_multidao_presenca"] = {nm: float(fp[k]) for k, nm in enumerate(names) if np.isfinite(fp[k])}
        out["spearman_km_vs_presente"] = m["presenca"]["spearman"]
        out["spearman_furto_km_vs_pedestre"] = m["furto_rua"]
    return out


# --- bootstrap ----------------------------------------------------------------------------------

def multiplicidades(dom_zona: np.ndarray, n_boot: int, seed: int = SEED) -> np.ndarray:
    """Multiplicidade de cada domicílio em cada réplica (B × n_dom): sorteio com reposição dentro da zona."""
    rng = np.random.default_rng(seed)
    ordem = np.argsort(dom_zona, kind="stable")             # domicílios agrupados por zona
    _, inicio, tamanho = np.unique(dom_zona[ordem], return_index=True, return_counts=True)
    ini, tam = np.repeat(inicio, tamanho), np.repeat(tamanho, tamanho)  # a zona de cada vaga
    M = np.zeros((n_boot, len(dom_zona)), dtype=np.int32)
    for b in range(n_boot):
        pick = ini + np.floor(rng.random(len(dom_zona)) * tam).astype(int)
        M[b] = np.bincount(ordem[pick], minlength=len(dom_zona))
    return M


def bootstrap(prep: dict, M: np.ndarray, g: dict, mp: dict, E, names, cr, log=print) -> pl.DataFrame:
    """Refaz a grade de rua com os pesos de cada réplica e recalcula (b), (c), (d) e o ranking."""
    rows = []
    kk = {nm: names.index(nm) for nm in CLASSES_TEXTO if nm in names}
    for b in range(M.shape[0]):
        for var, (cell, W, u) in prep.items():
            D, Rx = densidade(somar(cell, W * M[b, u][:, None], g), mp)
            m = metricas(E, names, D, Rx, cr)
            pp, rl = m["pessoa"], m["relogio"]
            rows.append({"replica": b, "variante": var, "spearman": pp["spearman"], "hora_pico": rl["pico"],
                         "multiplo_pico": rl["multiplo"], "hora_pico_pessoas": rl["pico_pessoas"],
                         "multiplo_pico_pessoas": rl["multiplo_pessoas"], "correlacao_mediana": rl["cor_mediana"],
                         "risco_madrugada": rl["risco_madrugada"],
                         **{f"risco_h{h:02d}": float(rl["risco"][h]) for h in range(24)},
                         **{f"subida_{k}": float(pp["subida"][i]) for k, i in kk.items()},
                         **{f"multidao_{nm}": float(pp["fator_multidao"][k]) for k, nm in enumerate(names)}})
        if (b + 1) % 10 == 0:
            log(f"bootstrap: {b + 1}/{M.shape[0]} réplicas")
    return pl.DataFrame(rows)


def _ic(v) -> list[float]:
    v = np.asarray(v, float)
    v = v[np.isfinite(v)]
    return [float(np.percentile(v, 2.5)), float(np.percentile(v, 97.5))] if len(v) else [np.nan, np.nan]


def resumo_boot(bt: pl.DataFrame, var: str) -> dict:
    d = bt.filter(pl.col("variante") == var)
    picos = d["hora_pico"].value_counts().sort("count", descending=True)
    out = {"B": d.height, "spearman_ic95": _ic(d["spearman"]), "multiplo_pico_ic95": _ic(d["multiplo_pico"]),
           "risco_h03_ic95": _ic(d["risco_h03"]), "correlacao_mediana_ic95": _ic(d["correlacao_mediana"]),
           "multiplo_pico_pessoas_ic95": _ic(d["multiplo_pico_pessoas"]), "risco_madrugada_ic95": _ic(d["risco_madrugada"]),
           "risco_h01_ic95": _ic(d["risco_h01"]),
           "hora_pico_freq": {str(int(h)): int(c) for h, c in zip(picos["hora_pico"], picos["count"])}}
    for nm in CLASSES_TEXTO:
        c = f"subida_{nm}"
        if c in d.columns:
            out[f"P_sobe_{nm}"] = float((d[c] > 0).mean())
            out[f"P_cai_{nm}"] = float((d[c] < 0).mean())
            out[f"subida_ic95_{nm}"] = _ic(d[c])
    return out


# --- 2017 de referência -------------------------------------------------------------------------

def referencia_2017(E, names, ids, cr, mp=None, setor_ids=None) -> dict:
    """Métricas de 2017 a partir da tabela por setor salva por ``presenca`` (mesmas funções de 2023)."""
    tab = pl.read_parquet(P.out_dir() / "setor_presenca_hora.parquet")
    D = {c: np.nan_to_num(P._densidade(ids, tab, c)) for c in ("presenca", "rua")}
    w = tab.pivot(on="hora", index="cd_setor", values="rua").fill_null(0.0)
    Rx = ids.select("cd_setor").join(w, on="cd_setor", how="left").fill_null(0.0).select([str(h) for h in range(24)]).to_numpy()
    out = _resumo_metricas(metricas(E, names, D["rua"], Rx, cr, Dpres=D["presenca"]), names)
    if mp is not None:  # a densidade deste módulo reproduz a de presenca._densidade
        R17 = pl.DataFrame({"cd_setor": setor_ids}).join(w, on="cd_setor", how="left").fill_null(0.0).select(
            [str(h) for h in range(24)]).to_numpy()
        out["densidade_max_dif"] = float(np.max(np.abs(densidade(R17, mp)[0] - D["rua"])))
    salvo = json.loads((P.out_dir() / "presenca.json").read_text())
    out["confere_com_presenca_json"] = {
        "spearman": salvo[ROUBO]["rua"]["spearman_km_vs_pessoa"], "correlacao_mediana": salvo["relogio_vs_rua"]["correlacao_mediana"],
        "multiplo_pico": max(salvo["relogio_vs_rua"]["risco_relativo_por_hora"]), "hora_pico": salvo["relogio_vs_rua"]["hora_maior_risco"]}
    return out


def verifica_grade_2017(g: dict, setores, log=print) -> dict:
    """Refaz a grade de 2017 com as funções deste módulo e compara com a tabela salva por ``presenca``."""
    od = read_od2017(P.COLS)
    od = od[od["MUNI_DOM"].notna()]
    pres, rua = segmentos(od, ATIVOS["pe_bici"])
    tab = pl.read_parquet(P.out_dir() / "setor_presenca_hora.parquet").sort(["cd_setor", "hora"])
    ids = setores["cd_setor"].astype(str).to_numpy()
    out = {}
    for nome, seg in (("presenca", pres), ("rua", rua)):
        cell, W, _ = celulas(seg, g)
        novo = pl.DataFrame({"cd_setor": np.repeat(ids, 24), "hora": np.tile(np.arange(24), len(ids)),
                             "v": somar(cell, W, g).ravel()}).sort(["cd_setor", "hora"])
        a, b = novo["v"].to_numpy(), tab[nome].to_numpy()
        out[nome] = {"max_dif_relativa": float(np.max(np.abs(a - b)) / max(np.max(np.abs(b)), 1e-12)),
                     "total_novo": float(a.sum()), "total_salvo": float(b.sum())}
    log(f"verificação 2017: diferença relativa máxima presença {out['presenca']['max_dif_relativa']:.1e}, "
        f"rua {out['rua']['max_dif_relativa']:.1e}")
    return out


# --- execução -----------------------------------------------------------------------------------

def _prepara_2023(g: dict) -> tuple[dict, np.ndarray]:
    """Segmentos da OD 2023 já na grade: presença e rua nas duas definições, com o domicílio de cada um."""
    od = read_od2023()
    od = od[od["MUNI_DOM"].notna()].copy()
    doms = od[["ID_DOM", "ZONA"]].drop_duplicates("ID_DOM").reset_index(drop=True)
    od["_u"] = od["ID_DOM"].map(pd.Series(doms.index, index=doms["ID_DOM"])).astype(int)
    prep = {"rua": {}}
    for var, at in ATIVOS.items():
        pres, rua = segmentos(od, at)
        prep["rua"][var] = celulas(rua, g)
    prep["presenca"] = celulas(pres, g)
    return prep, doms["ZONA"].to_numpy()


def linhas_relogio(var: str, names: list[str], m: dict, cr: dict) -> list[dict]:
    """Perfil horário do roubo e dos pedestres por classe (painel 4b), como ``relogio_vs_rua.parquet`` de 2017."""
    Pr = m["relogio"]["Pr"]
    return [{"variante": var, "classe": nm, "peso": float(cr["w"][k]), **{f"crime_h{h:02d}": float(cr["Pc"][k, h]) for h in range(24)},
             **{f"rua_h{h:02d}": float(Pr[k, h]) for h in range(24)}} for k, nm in enumerate(names)]


def estabilidade(r17: dict, r23: dict) -> dict:
    """Concordância entre 2017 e 2023: fatores de multidão, perfil horário dos pedestres e ranking por pedestre."""
    from scipy.stats import spearmanr

    out = {}
    for k in ("fator_multidao_rua", "fator_multidao_presenca"):
        cls = sorted(set(r17[k]) & set(r23[k]))
        out[f"spearman_{k}"] = float(spearmanr([r17[k][c] for c in cls], [r23[k][c] for c in cls]).statistic)
    out["pearson_perfil_rua_cidade"] = float(np.corrcoef(r17["perfil_rua_cidade"], r23["perfil_rua_cidade"])[0, 1])
    out["pearson_log_risco_por_hora"] = float(np.corrcoef(np.log(r17["risco_relativo_por_hora"]),
                                                          np.log(r23["risco_relativo_por_hora"]))[0, 1])
    s17, s23 = r17["subida_no_ranking"], r23["subida_no_ranking"]
    cls = sorted(set(s17) & set(s23))
    out["mesma_direcao_no_ranking"] = int(sum(np.sign(s17[c]) == np.sign(s23[c]) for c in cls))
    out["classes_no_ranking"] = len(cls)
    return out


def _hora_cidade(rep: dict, bt: pl.DataFrame, crime: np.ndarray) -> pl.DataFrame:
    out = {"hora": np.arange(24), "crime": crime, "risco_2017": rep["2017"]["risco_relativo_por_hora"]}
    for var in ATIVOS:
        d = bt.filter(pl.col("variante") == var)
        ic = np.array([_ic(d[f"risco_h{h:02d}"]) for h in range(24)])
        out |= {f"risco_{var}": rep["2023"][var]["risco_relativo_por_hora"], f"risco_{var}_ic_inf": ic[:, 0],
                f"risco_{var}_ic_sup": ic[:, 1], f"rua_{var}": rep["2023"][var]["perfil_rua_cidade"]}
    return pl.DataFrame(out)


def run(log=print, n_boot: int = N_BOOT) -> dict:
    import pyogrio

    from cityomics.atlas.fenotipo import setor_exposicao
    from cityomics.units.setores import load_setores

    zip_path = guard.pre2026(external_dir() / OD23_ZIP)
    sha = sha256(zip_path) if zip_path.exists() else None
    log(f"OD 2023: sha256 do zip {'confere' if sha == OD23_SHA256 else 'NÃO confere: ' + str(sha)}")
    cols23 = list(pyogrio.read_info(guard.pre2026(external_dir() / OD23_DBF))["fields"])
    cols17 = list(pyogrio.read_info(guard.pre2026(external_dir() / P.OD_DBF))["fields"])
    rep = {"fonte": {"zip": OD23_ZIP, "sha256": sha, "sha256_confere": sha == OD23_SHA256, "dbf": OD23_DBF},
           "anos_crime": list(P.ANOS), "dias": "úteis", "sigma_m": P.SIGMA_M, "res_m": P.RES}
    rep["compatibilidade"] = compatibilidade(read_od2017(), read_od2023(), cols23, cols17)
    c17, c23 = rep["compatibilidade"]["2017"], rep["compatibilidade"]["2023"]
    log(f"compatibilidade: viagens de residentes do município {c17['sp_viagens_expandidas']:,.0f} (2017) × "
        f"{c23['sp_viagens_expandidas']:,.0f} (2023); a pé {c17['sp_a_pe_expandidas']:,.0f} × {c23['sp_a_pe_expandidas']:,.0f}")

    setores = load_setores()
    setor_ids = setores["cd_setor"].astype(str).to_numpy()
    g = grade(setores)
    rep["verificacao_grade_2017"] = verifica_grade_2017(g, setores, log=log)

    prep, dom_zona = _prepara_2023(g)
    Pm = somar(*prep["presenca"][:2], g)
    Rm = {var: somar(cell, W, g) for var, (cell, W, _) in prep["rua"].items()}
    pl.DataFrame({"cd_setor": np.repeat(setor_ids, 24), "hora": np.tile(np.arange(24), len(setor_ids)), "presenca": Pm.ravel(),
                  **{f"rua_{v}": R.ravel() for v, R in Rm.items()}}).write_parquet(out_dir() / "setor_presenca_hora.parquet")
    t17 = pl.read_parquet(P.out_dir() / "setor_presenca_hora.parquet")
    rep["totais"] = {"2017": {"presenca_3h": float(t17.filter(pl.col("hora") == 3)["presenca"].sum()),
                              "presenca_15h": float(t17.filter(pl.col("hora") == 15)["presenca"].sum()),
                              "rua_pe_bici_pessoas_hora_dia": float(t17["rua"].sum())},
                     "2023": {"presenca_3h": float(Pm[:, 3].sum()), "presenca_15h": float(Pm[:, 15].sum()),
                              **{f"rua_{v}_pessoas_hora_dia": float(R.sum()) for v, R in Rm.items()},
                              **{f"rua_{v}_hora_pico": int(np.argmax(R.sum(0))) for v, R in Rm.items()}}}
    log(f"2023: presença às 3h {Pm[:, 3].sum():,.0f}; rua (a pé + bici) {Rm['pe_bici'].sum():,.0f} pessoas-hora/dia")

    ids, E, names = setor_exposicao()
    mp = mapa_setores(setor_ids, ids)
    cr = crime_cache(ids, E, names)
    rep["eventos"] = {"celular_roubo": cr["per"][ROUBO]["eventos"], "celular_furto": cr["per"]["celular_furto"]["eventos"],
                      "roubo_hora_exata": cr["eventos_hora"]}
    rep["2017"] = referencia_2017(E, names, ids, cr, mp=mp, setor_ids=setor_ids)
    Dpres, _ = densidade(Pm, mp)
    rep["2023"], rows, rel = {}, [], []
    for var, R in Rm.items():
        D, Rx = densidade(R, mp)
        m = metricas(E, names, D, Rx, cr, Dpres=Dpres)
        rep["2023"][var] = _resumo_metricas(m, names)
        rel += linhas_relogio(var, names, m, cr)
        for k, nm in enumerate(names):
            for j, q in enumerate(cr["periodos"]):
                rows.append({"variante": var, "classe": nm, "periodo": q, "lambda_km": float(cr["per"][ROUBO]["lam"][k, j]),
                             "rho_pessoa": float(m["pessoa"]["rho"][k, j]), "densidade_relativa": float(m["pessoa"]["dens_k"][k, j])})
        log(f"2023 {var}: ρ(km × pedestre) = {m['pessoa']['spearman']:.2f}; pico às {m['relogio']['pico']}h "
            f"({m['relogio']['multiplo']:.1f}×); correlação mediana = {m['relogio']['cor_mediana']:.2f}")
    pl.DataFrame(rows).write_parquet(out_dir() / "risco_por_pessoa.parquet")
    pl.DataFrame(rel).write_parquet(out_dir() / "relogio_vs_rua.parquet")
    rep["estabilidade_2017_2023"] = {v: estabilidade(rep["2017"], rep["2023"][v]) for v in ATIVOS}
    (out_dir() / "presenca_2023.json").write_text(json.dumps(rep, ensure_ascii=False, indent=2))

    M = multiplicidades(dom_zona, n_boot)
    bt = bootstrap(prep["rua"], M, g, mp, E, names, cr, log=log)
    bt.write_parquet(out_dir() / "bootstrap.parquet")
    for var in ATIVOS:
        rep["2023"][var]["bootstrap"] = resumo_boot(bt, var)
    rep["bootstrap"] = {"unidade": "domicílio (ID_DOM)", "estrato": "zona de residência (ZONA)", "B": int(n_boot), "semente": SEED,
                        "domicilios": int(len(dom_zona)), "zonas": int(len(np.unique(dom_zona)))}
    _hora_cidade(rep, bt, cr["cidade_crime"]).write_parquet(out_dir() / "risco_hora_cidade.parquet")
    (out_dir() / "presenca_2023.json").write_text(json.dumps(rep, ensure_ascii=False, indent=2))
    return rep
